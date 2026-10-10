"""Read lead spreadsheets into the existing background import format."""

import csv
import io
from pathlib import Path

from ninja.errors import HttpError

MAX_IMPORT_BYTES = 25 * 1024 * 1024
EXTENSIONS = {'.csv', '.xlsx', '.xls'}


def _normalise_header(value):
    return str(value or '').strip().lower().replace(' ', '_').replace('-', '_')


def parse_lead_file(upload):
    extension = Path(upload.name).suffix.lower()
    if extension not in EXTENSIONS:
        raise HttpError(400, 'Choose a CSV or Excel file (.csv, .xlsx, .xls).')
    if not upload.size or upload.size > MAX_IMPORT_BYTES:
        raise HttpError(400, 'The file must be larger than 0 bytes and no more than 25 MB.')

    try:
        if extension == '.csv':
            rows = csv.reader(io.TextIOWrapper(upload.file, encoding='utf-8-sig', newline=''))
            data = _rows_to_contacts(rows)
        elif extension == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(upload.file, read_only=True, data_only=True)
            try:
                data = _rows_to_contacts(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=upload.read())
            sheet = workbook.sheet_by_index(0)
            data = _rows_to_contacts(sheet.row_values(index) for index in range(sheet.nrows))
    except HttpError:
        raise
    except Exception as exc:
        # Spreadsheet parsers raise their own format-specific errors for corrupt files.
        raise HttpError(400, 'Could not read this CSV or Excel file. Check that it is valid.') from exc
    if not data:
        raise HttpError(400, 'No leads with an email address were found.')
    return data


def _rows_to_contacts(rows):
    iterator = iter(rows)
    headers = None
    for row in iterator:
        if any(str(value or '').strip() for value in row):
            headers = [_normalise_header(value) for value in row]
            break
    if not headers or 'email' not in headers:
        raise HttpError(400, 'The first row must contain an email column.')

    contacts = []
    for row in iterator:
        values = [str(value).strip() if value is not None else '' for value in row]
        record = {header: values[index] if index < len(values) else ''
                  for index, header in enumerate(headers) if header}
        if record.get('email'):
            contacts.append(record)
    return contacts
